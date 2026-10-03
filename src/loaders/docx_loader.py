from pathlib import Path

from docx import Document

from src.models import Chunk


def _is_heading(paragraph) -> bool:
    name = (paragraph.style.name or "").lower()
    return name.startswith("heading") or name.startswith("заголовок")


def load_docx(path: str | Path) -> list[Chunk]:
    path = Path(path)
    doc = Document(path)
    chunks: list[Chunk] = []
    title = "начало документа"
    buffer: list[str] = []

    def flush():
        text = "\n".join(buffer).strip()
        if text:
            chunks.append(Chunk(text=text, source=path.name, location=title))
        buffer.clear()

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if _is_heading(paragraph):
            flush()
            title = text
        else:
            buffer.append(text)
    flush()

    for i, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(rows).strip()
        if text:
            chunks.append(Chunk(text=text, source=path.name, location=f"таблица {i}"))

    return chunks
